"""
Builds the FAQ page from the client's spreadsheet.

    python tools/build-faq.py

Reads docs/content/FAQ.xlsx (one sheet per category: question in column A, answer in
column B) and writes:

  docs/content/faq.json   the questions and answers, as the site uses them
  faq/index.html          the category chips between <!-- @faq-nav --> markers, the
                          sections between <!-- @faq --> markers, and the FAQPage JSON-LD
                          between <!-- @faq-ld --> markers

The copy is the client's, verbatim, with whitespace tidied and the fixes in FIXES below.
Run tools/build.js afterwards, as for any page. Needs openpyxl.
"""
import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "content" / "FAQ.xlsx"
DATA = ROOT / "docs" / "content" / "faq.json"
PAGE = ROOT / "faq" / "index.html"
ORIGIN = "https://www.princelubricants.com"

# Sheet name: (section id, chip label, section label, heading with its gold part in <span>).
# The labels are functional: they name each range the way the answers describe it.
CATEGORIES = {
    "COMPANY FAQ": ("company", "Company", "The Company", "ABOUT PRINCE <span>LUBRICANTS</span>"),
    "TECHNOLOGY FAQ": ("technology", "Technology", "Ester &amp; Synthetic Technology", "OUR <span>TECHNOLOGY</span>"),
    "FS1 FAQ": ("fs1", "FS1 &amp; FS1 EUROGEN", "Passenger Car Motor Oils", "PRINCE FS1 <span>&amp; FS1 EUROGEN</span>"),
    "FSR FAQ": ("fsr", "FSR", "Racing Oils", "PRINCE <span>FSR</span>"),
    "FSe FAQ": ("fse", '<span class="keep-case">FSe</span>', "Resource Conserving Engine Oils", 'PRINCE <span class="keep-case">FSe</span>'),
    "HDEO FAQ": ("d-series", "D Series", "Heavy-Duty Diesel Engine Oils", "PRINCE <span>D SERIES</span>"),
    "MCO FAQ": ("maxx", "MAXX", "Motorcycle Engine Oils", "PRINCE <span>MAXX</span>"),
}

# Typos in the spreadsheet, corrected here so the fix survives a re-export.
FIXES = {
    "Why are esters used in high-performance engine oils": "Why are esters used in high-performance engine oils?",
}


def tidy(text):
    text = re.sub(r"\s+", " ", str(text)).strip()
    return FIXES.get(text, text)


def slug(text, taken):
    base = re.sub(r"[^a-z0-9]+", "-", text.lower().replace("+", " plus ")).strip("-")
    if len(base) > 64:
        base = base[:64].rsplit("-", 1)[0]
    name, n = base, 2
    while name in taken:
        name, n = f"{base}-{n}", n + 1
    taken.add(name)
    return name


def read():
    book = load_workbook(SOURCE, read_only=True)
    taken, data = set(), []
    for sheet in book.worksheets:
        if sheet.title not in CATEGORIES:
            raise SystemExit(f"unknown sheet {sheet.title!r}: add it to CATEGORIES")
        cid, chip, label, title = CATEGORIES[sheet.title]
        items = []
        for row in sheet.iter_rows(values_only=True):
            if not row or row[0] is None or row[1] is None:
                continue
            q, a = tidy(row[0]), tidy(row[1])
            items.append({"id": slug(q, taken), "q": q, "a": a})
        data.append({"id": cid, "chip": chip, "label": label, "title": title, "items": items})
    return data


def section(index, cat):
    tone = "bg-light" if index % 2 == 0 else "bg-off"
    count = len(cat["items"])
    rows = []
    for item in cat["items"]:
        rows.append(
            f'          <details class="faq-item" id="{item["id"]}">\n'
            f'            <summary><h3 class="faq-q">{html.escape(item["q"], quote=False)}</h3>'
            f'<span class="faq-icon" aria-hidden="true"></span></summary>\n'
            f'            <div class="faq-a"><p>{html.escape(item["a"], quote=False)}</p></div>\n'
            f"          </details>"
        )
    items = "\n".join(rows)
    dot = re.sub(r"<[^>]+>", "", cat["chip"]).replace("&amp;", "&")
    return f"""  <section class="section faq-cat {tone}" id="{cat["id"]}" data-light data-dot="{html.escape(dot)}">
    <div class="container">
      <div class="faq-grid">
        <div class="faq-head rv-l">
          <span class="disc-num" aria-hidden="true">{index + 1:02d}</span>
          <span class="sec-label">{cat["label"]}</span>
          <h2 class="sec-title sec-title--md">{cat["title"]}</h2>
          <div class="gold-bar"></div>
          <p class="faq-count">{count} questions</p>
        </div>
        <div class="faq-list rv">
{items}
        </div>
      </div>
    </div>
  </section>"""


def chips(data):
    rows = "\n".join(f'        <li><a href="#{c["id"]}">{c["chip"]}</a></li>' for c in data)
    return f"      <ul>\n{rows}\n      </ul>"


def json_ld(data):
    questions = [
        {"@type": "Question", "name": item["q"], "acceptedAnswer": {"@type": "Answer", "text": item["a"]}}
        for cat in data
        for item in cat["items"]
    ]
    graph = {
        "@context": "https://schema.org",
        "@graph": [
            {
                "@type": "FAQPage",
                "@id": f"{ORIGIN}/faq/#webpage",
                "url": f"{ORIGIN}/faq/",
                "name": "Frequently Asked Questions | PRINCE LUBRICANTS",
                "isPartOf": {"@id": f"{ORIGIN}/#website"},
                "about": {"@id": f"{ORIGIN}/#organization"},
                "mainEntity": questions,
            },
            {
                "@type": "BreadcrumbList",
                "itemListElement": [
                    {"@type": "ListItem", "position": 1, "name": "Home", "item": f"{ORIGIN}/"},
                    {"@type": "ListItem", "position": 2, "name": "FAQ", "item": f"{ORIGIN}/faq/"},
                ],
            },
        ],
    }
    body = json.dumps(graph, ensure_ascii=False, indent=2).replace("</", "<\\/")
    return f'<script type="application/ld+json">\n{body}\n</script>'


def stamp(page, name, body):
    pattern = re.compile(rf"<!-- @{name} -->.*?<!-- /@{name} -->", re.S)
    if not pattern.search(page):
        raise SystemExit(f"faq/index.html has no <!-- @{name} --> markers")
    return pattern.sub(lambda _: f"<!-- @{name} -->\n{body}\n<!-- /@{name} -->", page)


def main():
    data = read()
    DATA.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    page = PAGE.read_text(encoding="utf-8")
    page = stamp(page, "faq-ld", json_ld(data))
    page = stamp(page, "faq-nav", chips(data))
    page = stamp(page, "faq", "\n\n".join(section(i, c) for i, c in enumerate(data)))
    PAGE.write_text(page, encoding="utf-8", newline="\n")
    total = sum(len(c["items"]) for c in data)
    print(f"{total} questions in {len(data)} categories -> {PAGE.relative_to(ROOT)}, {DATA.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
